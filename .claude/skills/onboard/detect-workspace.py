#!/usr/bin/env python3
"""Step 0: detect workspace copies + bootstrap missing live files.

Two jobs, run together:
  1. DETECT  - find every copy of this workspace, flag cloud-synced / duplicate copies.
  2. BOOTSTRAP - deterministically create any missing live file from its blank:
       *.example -> live file, kk-post-template.md -> kk-post.md, and blank trackers.
     Never overwrites an existing file, so real data is safe.

Run with the host's NATIVE python (not a translation layer like WSL/Git-bash),
so cloud placeholders (iCloud/OneDrive) are enumerated.

  python .claude/skills/onboard/detect-workspace.py            # detect + bootstrap
  python .claude/skills/onboard/detect-workspace.py --detect-only

Exit 0 = single local copy, safe to proceed. Exit 1 = a problem needs a decision
(more than one copy, or a cloud-synced working copy). Bootstrap of the current
copy still runs either way.
"""
import os, sys, shutil, time

MARKERS = ("CLAUDE.md", "config.json", "START-HERE.md")  # all three => a workspace copy
CLOUD = ("icloud", "onedrive", "dropbox", "google drive", "mobile documents")
SKIP_DIRS = {".git", "node_modules", ".obsidian", "__pycache__", "AppData", ".cache"}
MAX_DEPTH = 4

# Tracker schemas (kept in sync with onboard Step 0 0b).
TRACKERS = {
    "lead-gen-tracker.xlsx": ("Prospects", [
        "Name", "Country", "Job Title", "Company", "Status", "Date Requested",
        "Date Connected", "DM Stage", "Last Contact", "Next Action",
        "Touchpoints", "Lead Magnet Sent", "LinkedIn URL", "Notes"]),
    "engagement-tracker.xlsx": ("Accounts", [
        "Name", "Country", "Category", "Last Engaged", "Next Action",
        "Total Engagements", "First Engaged", "LinkedIn URL", "Date Added", "Notes"]),
}


def is_workspace(d):
    return all(os.path.exists(os.path.join(d, m)) for m in MARKERS)


def cloud_hits(path):
    low = path.lower()
    return [c for c in CLOUD if c in low]


def find_copies(home):
    roots = [
        home,
        os.path.join(home, "Desktop"), os.path.join(home, "Documents"),
        os.path.join(home, "Downloads"), os.path.join(home, "OneDrive"),
        os.path.join(home, "iCloudDrive"), os.path.join(home, "Dropbox"),
        os.path.join(home, "Google Drive"),
        os.path.join(home, "Library", "Mobile Documents", "com~apple~CloudDocs"),
    ]
    seen = set()
    roots = [r for r in roots if not (r in seen or seen.add(r))]
    found = set()
    for root in roots:
        if not os.path.isdir(root):
            continue
        for dirpath, dirnames, _ in os.walk(root):
            depth = dirpath[len(root):].count(os.sep)
            if depth >= MAX_DEPTH:
                dirnames[:] = []
                continue
            dirnames[:] = [d for d in dirnames if d not in SKIP_DIRS]
            if is_workspace(dirpath):
                found.add(os.path.realpath(dirpath))
                dirnames[:] = []  # don't descend into a found workspace
    return sorted(found)


def bootstrap(root):
    """Create any missing live file from its blank. Never overwrites. Returns (created, skipped_notes)."""
    created, notes = [], []

    # 1. every *.example -> its live file (account-profile.md.example -> account-profile.md, etc.)
    #    .env is deliberately excluded: it holds secrets/API keys and must be set up
    #    consciously (onboard Step 7), never auto-materialised blank.
    for dirpath, dirnames, filenames in os.walk(root):
        dirnames[:] = [d for d in dirnames if d not in SKIP_DIRS]
        for fn in filenames:
            if fn.endswith(".example"):
                live = os.path.join(dirpath, fn[: -len(".example")])
                base = os.path.basename(live)
                if base == ".env":
                    if not os.path.exists(live):
                        notes.append(".env not auto-created (secrets). Set API keys via onboard Step 7 if needed.")
                    continue
                if not os.path.exists(live):
                    try:
                        shutil.copyfile(os.path.join(dirpath, fn), live)
                        created.append(os.path.relpath(live, root))
                    except OSError as e:
                        notes.append(f"could not create {os.path.relpath(live, root)}: {e}")

    # 2. kk-post.md is special: its blank is kk-post-template.md (not a *.example file)
    kk, tmpl = os.path.join(root, "kk-post.md"), os.path.join(root, "kk-post-template.md")
    if not os.path.exists(kk) and os.path.exists(tmpl):
        try:
            shutil.copyfile(tmpl, kk)
            created.append("kk-post.md")
        except OSError as e:
            notes.append(f"could not create kk-post.md: {e}")

    # 3. blank trackers with their schemas (needs openpyxl)
    tdir = os.path.join(root, "trackers")
    need = [f for f in TRACKERS if not os.path.exists(os.path.join(tdir, f))]
    if need:
        try:
            import openpyxl
        except ImportError:
            notes.append("openpyxl not installed; trackers NOT created. "
                         "Run /onboard or `pip install openpyxl`, then re-run.")
        else:
            os.makedirs(tdir, exist_ok=True)
            for f in need:
                sheet, headers = TRACKERS[f]
                try:
                    wb = openpyxl.Workbook()
                    ws = wb.active
                    ws.title = sheet
                    ws.append(headers)
                    wb.save(os.path.join(tdir, f))
                    created.append(f"trackers/{f}")
                except OSError as e:
                    notes.append(f"could not create trackers/{f}: {e}")

    return created, notes


def main():
    detect_only = "--detect-only" in sys.argv
    home = os.path.expanduser("~")
    cwd = os.path.realpath(os.getcwd())

    copies = find_copies(home)
    print(f"Working copy: {cwd}")
    if cloud_hits(cwd):
        print(f"  WARNING: cloud-synced location ({', '.join(cloud_hits(cwd))}). "
              "Files may be online-only placeholders the tooling can't read/write. "
              'Set the folder to "Always keep on this device" or run from a local copy.')

    print(f"\nWorkspace copies found: {len(copies)}")
    for c in copies:
        try:
            mt = time.strftime("%Y-%m-%d %H:%M", time.localtime(os.path.getmtime(os.path.join(c, "config.json"))))
        except OSError:
            mt = "??"
        tags = ""
        if cloud_hits(c):
            tags += " [cloud]"
        if c == cwd:
            tags += " <-- CURRENT (Claude writes here)"
        print(f"  {c}  (config.json {mt}){tags}")

    if len(copies) > 1:
        print("\nACTION REQUIRED: more than one copy exists. They will drift, because Claude writes to the "
              "current copy while you may open another. Pick ONE source of truth before continuing.")

    if not detect_only:
        created, notes = bootstrap(cwd)
        print(f"\nBootstrap (current copy): created {len(created)} missing file(s).")
        for c in created:
            print(f"  + {c}")
        for n in notes:
            print(f"  ! {n}")
        if not created and not notes:
            print("  (nothing missing)")
        if len(copies) > 1:
            print("  NOTE: bootstrapped the CURRENT copy only. If it is not your source of truth, resolve the copies above.")

    rc = 0
    if len(copies) > 1 or cloud_hits(cwd):
        rc = 1
    else:
        print("\nOK: single local copy.")
    return rc


if __name__ == "__main__":
    sys.exit(main())
